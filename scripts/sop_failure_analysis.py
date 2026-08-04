"""SOP Test Plan Failure Analysis.

Analyzes failed tests from the SOP Test Plans referenced on Confluence page
"Test Plan SOP - Workstream Applications" (MLBEVO/2378907792) and correlates
them with their KPM tickets (MLBEVO Bugs labelled ``Pdig_kpm_bugs`` whose
summary starts with ``[KPM_bug_NNNNNNNN]``).

Correlation strategy (in order):

1. Direct Jira issue links from the failed Test to any Bug in the MLBEVO
   project (rarely used but included when present).
2. Xray Test Run defects (queried via ``/rest/raven/2.0/api/testruns``) — this
   is where Xray natively stores per-run defect links.
3. Reference list of all MLBEVO KPM bugs is written to a dedicated sheet so
   that the user can manually cross-match when automatic linkage does not
   exist.

Output: sop_failure_analysis.xlsx in the workspace root.

Usage (from mcporsche repo root, using the jira-mcp venv):
    jira-mcp\\venv\\Scripts\\python.exe scripts\\sop_failure_analysis.py
"""
from __future__ import annotations

import os
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --- Configuration ----------------------------------------------------------

# Load JIRA_PAT from jira-mcp/.env
ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / "jira-mcp" / ".env")

JIRA_BASE_URL = os.getenv("JIRA_BASE_URL", "https://api.skyway.porsche.com/jira").rstrip("/")
JIRA_PAT = os.getenv("JIRA_PAT")
if not JIRA_PAT:
    print("ERROR: JIRA_PAT not set in jira-mcp/.env", file=sys.stderr)
    sys.exit(1)

PROXIES = {}
if os.getenv("HTTP_PROXY"):
    PROXIES["http"] = os.getenv("HTTP_PROXY")
if os.getenv("HTTPS_PROXY"):
    PROXIES["https"] = os.getenv("HTTPS_PROXY")
PROXIES = PROXIES or None

# SOP Test Plans extracted from the Confluence page's testPlanTests(...) references.
SOP_TEST_PLANS: list[str] = [
    "MLBEVO-11921", "MLBEVO-17687", "MLBEVO-17695", "MLBEVO-17696",
    "MLBEVO-17697", "MLBEVO-17698", "MLBEVO-17699", "MLBEVO-17700",
    "MLBEVO-17701", "MLBEVO-17702", "MLBEVO-17703", "MLBEVO-17794",
    "MLBEVO-17795", "MLBEVO-17796", "MLBEVO-17797", "MLBEVO-17798",
    "MLBEVO-17799", "MLBEVO-17800", "MLBEVO-17801", "MLBEVO-17802",
    "MLBEVO-17803", "MLBEVO-17804", "MLBEVO-17805", "MLBEVO-17806",
    "MLBEVO-17807", "MLBEVO-17808", "MLBEVO-17817", "MLBEVO-17818",
    "MLBEVO-17819", "MLBEVO-17820", "MLBEVO-17821", "MLBEVO-17823",
    "MLBEVO-17824", "MLBEVO-17825", "MLBEVO-17826", "MLBEVO-17827",
    "MLBEVO-17838", "MLBEVO-17844", "MLBEVO-17849", "MLBEVO-17852",
    "MLBEVO-17853", "MLBEVO-17854", "MLBEVO-17855",
]

OUTPUT_XLSX = ROOT / "sop_failure_analysis.xlsx"

KPM_SUMMARY_RE = re.compile(r"\[KPM_bug_(\d+)\]", re.IGNORECASE)

# --- HTTP session -----------------------------------------------------------

session = requests.Session()
session.headers.update({
    "Authorization": f"Bearer {JIRA_PAT}",
    "Accept": "application/json",
})
if PROXIES:
    session.proxies.update(PROXIES)


def jira_get(path: str, params: dict | None = None) -> dict:
    url = f"{JIRA_BASE_URL}{path}"
    r = session.get(url, params=params, timeout=60)
    r.raise_for_status()
    return r.json()


def jira_search(jql: str, fields: list[str], max_results: int = 200) -> list[dict]:
    """Paginated search returning issues with the given fields."""
    issues: list[dict] = []
    start_at = 0
    while True:
        payload = jira_get(
            "/rest/api/2/search",
            params={
                "jql": jql,
                "startAt": start_at,
                "maxResults": max_results,
                "fields": ",".join(fields),
            },
        )
        batch = payload.get("issues", [])
        issues.extend(batch)
        total = payload.get("total", 0)
        start_at += len(batch)
        if not batch or start_at >= total:
            break
    return issues


def xray_test_runs(test_plan_key: str, page_size: int = 200) -> list[dict]:
    """Return all Xray test runs for a test plan (paginated)."""
    runs: list[dict] = []
    page = 1
    while True:
        batch = jira_get(
            "/rest/raven/2.0/api/testruns",
            params={
                "testPlanKey": test_plan_key,
                "page": page,
                "limit": page_size,
            },
        )
        if not isinstance(batch, list) or not batch:
            break
        runs.extend(batch)
        if len(batch) < page_size:
            break
        page += 1
    return runs


# --- Data model -------------------------------------------------------------

@dataclass
class TestPlan:
    key: str
    summary: str
    status: str


@dataclass
class KpmLink:
    bug_key: str
    kpm_id: str
    summary: str
    status: str
    priority: str


@dataclass
class FailedTest:
    test_plan_key: str
    test_plan_summary: str
    test_key: str
    test_summary: str
    test_status: str
    test_exec_keys: list[str]  # Xray executions where this test currently fails
    kpms: list[KpmLink]        # linked KPMs (via issue links + Xray run defects)


@dataclass
class KpmBug:
    key: str
    kpm_id: str
    summary: str
    status: str
    priority: str
    fix_versions: list[str]
    labels: list[str]
    created: str
    updated: str
    epic: str
    reporter: str
    assignee: str


# --- Helpers ----------------------------------------------------------------

def parse_kpm(bug_summary: str) -> str | None:
    m = KPM_SUMMARY_RE.search(bug_summary or "")
    return m.group(1) if m else None


def extract_linked_bugs(issue_fields: dict) -> list[dict]:
    """Return the linked issue dicts that are bugs referenced from a Test."""
    bugs = []
    for link in issue_fields.get("issuelinks", []) or []:
        for direction in ("outwardIssue", "inwardIssue"):
            linked = link.get(direction)
            if not linked:
                continue
            itype = (linked.get("fields", {}).get("issuetype") or {}).get("name", "")
            if itype.lower() == "bug":
                bugs.append(linked)
    return bugs


# --- Main -------------------------------------------------------------------

def fetch_test_plans() -> dict[str, TestPlan]:
    print(f"[1/4] Fetching {len(SOP_TEST_PLANS)} SOP test plan headers...")
    jql = f"key in ({','.join(SOP_TEST_PLANS)})"
    issues = jira_search(jql, fields=["summary", "status"])
    plans: dict[str, TestPlan] = {}
    for i in issues:
        f = i["fields"]
        plans[i["key"]] = TestPlan(
            key=i["key"],
            summary=f.get("summary", ""),
            status=(f.get("status") or {}).get("name", ""),
        )
    print(f"      -> got {len(plans)} plans")
    return plans


def fetch_failed_tests_for_plan(plan_key: str) -> list[dict]:
    jql = (
        f'issue in testPlanTests({plan_key}) '
        f'and testrunstatus = "{plan_key}- fail"'
    )
    return jira_search(
        jql,
        fields=["summary", "status", "issuelinks"],
        max_results=100,
    )


def fetch_xray_defects_for_plan(plan_key: str) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """Return ({testKey: [defectKey,...]}, {testKey: [testExecKey,...]}) for FAIL runs."""
    try:
        runs = xray_test_runs(plan_key)
    except requests.HTTPError as e:
        print(f"  ! Xray runs {plan_key}: {e}", file=sys.stderr)
        return {}, {}
    defects_by_test: dict[str, list[str]] = {}
    execs_by_test: dict[str, list[str]] = {}
    for run in runs:
        if (run.get("status") or "").upper() != "FAIL":
            continue
        test_key = run.get("testKey")
        exec_key = run.get("testExecKey")
        if not test_key:
            continue
        if exec_key:
            execs_by_test.setdefault(test_key, []).append(exec_key)
        for d in run.get("defects") or []:
            # Xray v2 returns defects as strings (issue keys) or objects
            if isinstance(d, str):
                defects_by_test.setdefault(test_key, []).append(d)
            elif isinstance(d, dict):
                key = d.get("key") or d.get("issueKey")
                if key:
                    defects_by_test.setdefault(test_key, []).append(key)
    return defects_by_test, execs_by_test


def fetch_bugs_by_keys(keys: Iterable[str]) -> dict[str, dict]:
    keys = list({k for k in keys if k})
    if not keys:
        return {}
    result: dict[str, dict] = {}
    chunk = 100
    for i in range(0, len(keys), chunk):
        subset = keys[i : i + chunk]
        jql = f"key in ({','.join(subset)})"
        issues = jira_search(jql, fields=["summary", "status", "priority"])
        for it in issues:
            result[it["key"]] = it
    return result


def fetch_all_kpm_bugs() -> list[KpmBug]:
    """All MLBEVO KPM bugs (label Pdig_kpm_bugs)."""
    print("[3/4] Fetching all MLBEVO KPM bugs (label = Pdig_kpm_bugs)...")
    issues = jira_search(
        'project = MLBEVO AND labels = "Pdig_kpm_bugs" ORDER BY created DESC',
        fields=[
            "summary", "status", "priority", "fixVersions", "labels",
            "created", "updated", "customfield_10008", "reporter", "assignee",
        ],
        max_results=200,
    )
    out: list[KpmBug] = []
    for i in issues:
        f = i["fields"]
        kpm_id = parse_kpm(f.get("summary", "")) or ""
        out.append(KpmBug(
            key=i["key"],
            kpm_id=kpm_id,
            summary=f.get("summary", ""),
            status=(f.get("status") or {}).get("name", ""),
            priority=(f.get("priority") or {}).get("name", ""),
            fix_versions=[fv.get("name", "") for fv in f.get("fixVersions") or []],
            labels=f.get("labels") or [],
            created=(f.get("created") or "")[:10],
            updated=(f.get("updated") or "")[:10],
            epic=f.get("customfield_10008") or "",
            reporter=(f.get("reporter") or {}).get("displayName", ""),
            assignee=(f.get("assignee") or {}).get("displayName", ""),
        ))
    print(f"      -> got {len(out)} KPM bug(s)")
    return out


def build_report() -> tuple[list[FailedTest], list[KpmBug]]:
    plans = fetch_test_plans()

    print("[2/4] Querying failed tests + Xray runs for each SOP test plan...")
    all_failed: list[tuple[str, dict]] = []      # (plan_key, test_issue)
    xray_defects: dict[str, dict[str, list[str]]] = {}
    xray_execs:   dict[str, dict[str, list[str]]] = {}

    for idx, plan_key in enumerate(SOP_TEST_PLANS, 1):
        try:
            tests = fetch_failed_tests_for_plan(plan_key)
        except requests.HTTPError as e:
            print(f"  ! {plan_key}: {e}", file=sys.stderr)
            continue
        defects_by_test, execs_by_test = fetch_xray_defects_for_plan(plan_key)
        xray_defects[plan_key] = defects_by_test
        xray_execs[plan_key] = execs_by_test
        print(
            f"  ({idx}/{len(SOP_TEST_PLANS)}) {plan_key}: "
            f"{len(tests)} failed tests, "
            f"{sum(len(v) for v in defects_by_test.values())} Xray defect refs"
        )
        for t in tests:
            all_failed.append((plan_key, t))
        time.sleep(0.05)

    # collect bug keys from BOTH direct issue links and Xray run defects
    bug_keys: set[str] = set()
    for _, test in all_failed:
        for b in extract_linked_bugs(test["fields"]):
            bug_keys.add(b["key"])
    for plan_defects in xray_defects.values():
        for defs in plan_defects.values():
            bug_keys.update(defs)

    print(f"[3a/4] Resolving {len(bug_keys)} defect/bug candidate(s)...")
    bugs = fetch_bugs_by_keys(bug_keys)

    kpm_bugs = fetch_all_kpm_bugs()

    print("[4/4] Assembling report rows...")
    rows: list[FailedTest] = []
    for plan_key, test in all_failed:
        plan = plans.get(plan_key)
        test_key = test["key"]

        candidate_bug_keys: set[str] = set()
        for b in extract_linked_bugs(test["fields"]):
            candidate_bug_keys.add(b["key"])
        candidate_bug_keys.update(xray_defects.get(plan_key, {}).get(test_key, []))

        kpms: list[KpmLink] = []
        for bk in candidate_bug_keys:
            bug = bugs.get(bk)
            if not bug:
                continue
            bf = bug["fields"]
            summary = bf.get("summary", "")
            kpm_id = parse_kpm(summary) or ""
            kpms.append(KpmLink(
                bug_key=bk,
                kpm_id=kpm_id,
                summary=summary,
                status=(bf.get("status") or {}).get("name", ""),
                priority=(bf.get("priority") or {}).get("name", ""),
            ))

        rows.append(FailedTest(
            test_plan_key=plan_key,
            test_plan_summary=plan.summary if plan else "",
            test_key=test_key,
            test_summary=test["fields"].get("summary", ""),
            test_status=(test["fields"].get("status") or {}).get("name", ""),
            test_exec_keys=xray_execs.get(plan_key, {}).get(test_key, []),
            kpms=kpms,
        ))
    return rows, kpm_bugs


# --- Excel output -----------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FAIL_FILL = PatternFill("solid", fgColor="F8CBAD")

HEADERS = [
    "Test Plan", "Test Plan Summary",
    "Failed Test", "Test Summary", "Test Status",
    "Failing Test Executions",
    "KPM Bug (Jira)", "KPM ID", "KPM Summary", "KPM Status", "KPM Priority",
    "Jira Test Link", "KPM Bug Link",
]


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 0
        for cell in col:
            v = cell.value
            if v is None:
                continue
            width = max(width, min(len(str(v)), max_width))
        ws.column_dimensions[letter].width = width + 2


def write_excel(rows: list[FailedTest], kpm_bugs: list[KpmBug]) -> None:
    wb = Workbook()

    # --- Sheet 1: Failures (one row per failed test x KPM) -----------------
    ws = wb.active
    ws.title = "Failures"
    ws.append(HEADERS)
    for c, _ in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        test_url = f"{JIRA_BASE_URL}/browse/{row.test_key}"
        exec_str = ", ".join(row.test_exec_keys)
        if not row.kpms:
            ws.append([
                row.test_plan_key, row.test_plan_summary,
                row.test_key, row.test_summary, row.test_status,
                exec_str,
                "", "", "", "", "",
                test_url, "",
            ])
        else:
            for kpm in row.kpms:
                bug_url = f"{JIRA_BASE_URL}/browse/{kpm.bug_key}"
                ws.append([
                    row.test_plan_key, row.test_plan_summary,
                    row.test_key, row.test_summary, row.test_status,
                    exec_str,
                    kpm.bug_key, kpm.kpm_id, kpm.summary, kpm.status, kpm.priority,
                    test_url, bug_url,
                ])
    ws.freeze_panes = "A2"
    _autosize(ws)

    # --- Sheet 2: Summary per Test Plan ------------------------------------
    summary_ws = wb.create_sheet("Summary")
    summary_headers = [
        "Test Plan", "Test Plan Summary", "Failed Tests",
        "Failed Tests with KPM", "Failed Tests without KPM", "Unique KPM Bugs",
    ]
    summary_ws.append(summary_headers)
    for c, _ in enumerate(summary_headers, 1):
        cell = summary_ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    by_plan: dict[str, list[FailedTest]] = {}
    for r in rows:
        by_plan.setdefault(r.test_plan_key, []).append(r)

    for plan_key in SOP_TEST_PLANS:
        plan_rows = by_plan.get(plan_key, [])
        with_kpm = sum(1 for r in plan_rows if r.kpms)
        without_kpm = len(plan_rows) - with_kpm
        unique_kpms = {k.bug_key for r in plan_rows for k in r.kpms}
        plan_summary = plan_rows[0].test_plan_summary if plan_rows else ""
        summary_ws.append([
            plan_key, plan_summary, len(plan_rows),
            with_kpm, without_kpm, len(unique_kpms),
        ])
    summary_ws.freeze_panes = "A2"
    _autosize(summary_ws)

    # --- Sheet 3: Linked KPMs (only those tied to failed tests) ------------
    linked_ws = wb.create_sheet("Linked KPMs")
    linked_headers = ["KPM Bug", "KPM ID", "Summary", "Status", "Priority",
                      "Referenced By (# failed tests)", "Jira Link"]
    linked_ws.append(linked_headers)
    for c, _ in enumerate(linked_headers, 1):
        cell = linked_ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    kpm_index: dict[str, tuple[KpmLink, int]] = {}
    for r in rows:
        for kpm in r.kpms:
            existing = kpm_index.get(kpm.bug_key)
            if existing is None:
                kpm_index[kpm.bug_key] = (kpm, 1)
            else:
                kpm_index[kpm.bug_key] = (kpm, existing[1] + 1)

    for bug_key, (kpm, count) in sorted(kpm_index.items()):
        linked_ws.append([
            bug_key, kpm.kpm_id, kpm.summary, kpm.status, kpm.priority,
            count, f"{JIRA_BASE_URL}/browse/{bug_key}",
        ])
    linked_ws.freeze_panes = "A2"
    _autosize(linked_ws)

    # --- Sheet 4: All MLBEVO KPM Bugs (reference for manual matching) ------
    kpm_ws = wb.create_sheet("All KPM Bugs (MLBEVO)")
    kpm_headers = ["Jira Key", "KPM ID", "Summary", "Status", "Priority",
                   "Fix Versions", "Labels", "Created", "Updated",
                   "Reporter", "Assignee", "Epic Link", "Jira Link"]
    kpm_ws.append(kpm_headers)
    for c, _ in enumerate(kpm_headers, 1):
        cell = kpm_ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for b in kpm_bugs:
        kpm_ws.append([
            b.key, b.kpm_id, b.summary, b.status, b.priority,
            ", ".join(b.fix_versions), ", ".join(b.labels),
            b.created, b.updated, b.reporter, b.assignee, b.epic,
            f"{JIRA_BASE_URL}/browse/{b.key}",
        ])
    kpm_ws.freeze_panes = "A2"
    _autosize(kpm_ws)

    # --- Sheet 5: Notes ----------------------------------------------------
    notes_ws = wb.create_sheet("Notes")
    notes = [
        ["Source"],
        ["Confluence page", "Test Plan SOP - Workstream Applications"],
        ["Page URL", "https://skyway.porsche.com/confluence/spaces/MLBEVO/pages/2378907792"],
        [],
        ["Scope"],
        ["SOP Test Plans (Jira)", str(len(SOP_TEST_PLANS))],
        ["Failed test rows",     str(len(rows))],
        ["Rows with KPM link",   str(sum(1 for r in rows if r.kpms))],
        ["Unique KPM bugs",      str(len(kpm_index))],
        ["Total KPM bugs listed (label=Pdig_kpm_bugs)", str(len(kpm_bugs))],
        [],
        ["Method"],
        ["1. SOP test plans extracted from testPlanTests(...) references on the Confluence page (43 plans)."],
        ["2. Failed tests per plan queried via JQL:  issue in testPlanTests(<KEY>) and testrunstatus = \"<KEY>- fail\""],
        ["3. KPM candidates gathered from two sources:"],
        ["     a. Jira issue links (Bug type) on the failed test."],
        ["     b. Xray Test Run defects (GET /rest/raven/2.0/api/testruns?testPlanKey=<KEY>) filtered to FAIL runs."],
        ["4. Bug keys were resolved to summaries; the KPM ID is parsed from the [KPM_bug_NNNNNNNN] prefix."],
        [],
        ["Caveats"],
        ["- The Porsche Xray setup does NOT systematically populate the per-run 'defects' field."],
        ["- Consequently, most failed tests have no automated KPM linkage. Use the 'All KPM Bugs (MLBEVO)' sheet"],
        ["  as a reference to correlate manually (e.g. by component, fix version, or KPM ID)."],
        ["- 'Failing Test Executions' lists the Test Execution issues where the test currently reports FAIL,"],
        ["  which can help locate the correct execution to inspect in the Jira Xray UI."],
    ]
    for line in notes:
        notes_ws.append(line)
    for c in range(1, 3):
        notes_ws.cell(row=1, column=c).font = Font(bold=True)
    for header_row in (1, 5, 12, 20):
        cell = notes_ws.cell(row=header_row, column=1)
        cell.font = Font(bold=True, size=12)
    notes_ws.column_dimensions["A"].width = 45
    notes_ws.column_dimensions["B"].width = 90

    wb.save(OUTPUT_XLSX)
    print(f"\nWrote {OUTPUT_XLSX}")
    print(f"  Failed test rows    : {len(rows)}")
    print(f"  Rows with KPM link  : {sum(1 for r in rows if r.kpms)}")
    print(f"  Unique linked KPMs  : {len(kpm_index)}")
    print(f"  All MLBEVO KPM bugs : {len(kpm_bugs)}")


# --- Entry ------------------------------------------------------------------

if __name__ == "__main__":
    rows, kpm_bugs = build_report()
    write_excel(rows, kpm_bugs)
